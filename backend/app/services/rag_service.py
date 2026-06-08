from __future__ import annotations

import hashlib
import io
import math
import re
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, ClassVar, Protocol

from app.core.config import settings

try:
    from qdrant_client import QdrantClient
    from qdrant_client.http import models as qdrant_models
except ImportError:  # pragma: no cover - dependency is optional at import time
    QdrantClient = None
    qdrant_models = None

try:
    from fastembed import TextEmbedding
except ImportError:  # pragma: no cover - dependency is optional at import time
    TextEmbedding = None

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover - dependency is optional at import time
    PdfReader = None

try:
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover - dependency is optional at import time
    DocxDocument = None


@dataclass(frozen=True)
class ParsedDocument:
    document_id: str
    filename: str
    content_type: str | None
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class DocumentChunk:
    chunk_id: str
    document_id: str
    text: str
    metadata: dict[str, Any]


@dataclass(frozen=True)
class RetrievalResult:
    chunk_id: str
    document_id: str
    text: str
    score: float
    metadata: dict[str, Any]


class EmbeddingProvider(Protocol):
    dimension: int
    provider_name: str

    def embed(self, text: str) -> list[float]: ...

    def embed_many(self, texts: list[str]) -> list[list[float]]: ...


class DocumentParser:
    def parse(
        self,
        content: bytes,
        filename: str,
        content_type: str | None = None,
        metadata: dict[str, Any] | None = None,
    ) -> ParsedDocument:
        suffix = Path(filename).suffix.lower()
        parser_name = "text"
        if suffix == ".pdf" or content_type == "application/pdf":
            text = self._parse_pdf(content)
            parser_name = "pdf"
        elif suffix == ".docx" or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            text = self._parse_docx(content)
            parser_name = "docx"
        else:
            text = self._parse_text(content)

        normalized_text = self._normalize_text(text)
        return ParsedDocument(
            document_id=str(uuid.uuid4()),
            filename=filename,
            content_type=content_type,
            text=normalized_text,
            metadata={
                "filename": filename,
                "content_type": content_type,
                "byte_size": len(content),
                "parser": parser_name,
                "char_count": len(normalized_text),
                "created_at": datetime.now(timezone.utc).isoformat(),
                **(metadata or {}),
            },
        )

    def _parse_text(self, content: bytes) -> str:
        for encoding in ("utf-8", "utf-8-sig", "gb18030"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="ignore")

    def _parse_pdf(self, content: bytes) -> str:
        if PdfReader is None:
            raise ValueError("PDF 解析依赖 pypdf 未安装。")
        try:
            reader = PdfReader(io.BytesIO(content))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            raise ValueError("PDF 文件解析失败，请确认文件未加密且内容可复制。") from exc

    def _parse_docx(self, content: bytes) -> str:
        if DocxDocument is None:
            raise ValueError("DOCX 解析依赖 python-docx 未安装。")
        try:
            document = DocxDocument(io.BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            return "\n".join([*paragraphs, *table_cells])
        except Exception as exc:
            raise ValueError("DOCX 文件解析失败，请确认文件格式正确。") from exc

    def _normalize_text(self, text: str) -> str:
        text = re.sub(r"\r\n?", "\n", text)
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()


class TextChunker:
    def __init__(self, chunk_size: int = 800, overlap: int = 120) -> None:
        if chunk_size <= 0:
            raise ValueError("chunk_size must be greater than 0")
        if overlap < 0 or overlap >= chunk_size:
            raise ValueError("overlap must be non-negative and smaller than chunk_size")
        self.chunk_size = chunk_size
        self.overlap = overlap

    def split(self, document: ParsedDocument) -> list[DocumentChunk]:
        if not document.text:
            return []

        chunks: list[DocumentChunk] = []
        start = 0
        chunk_index = 0
        while start < len(document.text):
            end = min(start + self.chunk_size, len(document.text))
            chunk_text = document.text[start:end].strip()
            if chunk_text:
                chunks.append(
                    DocumentChunk(
                        chunk_id=f"{document.document_id}:{chunk_index}",
                        document_id=document.document_id,
                        text=chunk_text,
                        metadata={
                            **document.metadata,
                            "chunk_index": chunk_index,
                            "start_char": start,
                            "end_char": end,
                        },
                    )
                )
                chunk_index += 1
            if end == len(document.text):
                break
            start = max(end - self.overlap, start + 1)
        return chunks


class HashingEmbeddingProvider:
    provider_name = "hashing"

    def __init__(self, dimension: int = 384) -> None:
        self.dimension = dimension

    def embed_many(self, texts: list[str]) -> list[list[float]]:
        return [self.embed(text) for text in texts]

    def embed(self, text: str) -> list[float]:
        vector = [0.0] * self.dimension
        tokens = re.findall(r"[\w\u4e00-\u9fff]+", text.lower())
        for token in tokens:
            digest = hashlib.sha256(token.encode("utf-8")).digest()
            index = int.from_bytes(digest[:4], "big") % self.dimension
            sign = 1.0 if digest[4] % 2 == 0 else -1.0
            vector[index] += sign

        norm = math.sqrt(sum(value * value for value in vector))
        if norm == 0:
            return vector
        return [value / norm for value in vector]


class FastEmbedEmbeddingProvider:
    provider_name = "fastembed"

    def __init__(self, model_name: str) -> None:
        if TextEmbedding is None:
            raise RuntimeError("fastembed is not installed")
        self.model_name = model_name
        self._model = TextEmbedding(model_name=model_name)
        self.dimension = len(self.embed("dimension probe"))

    def embed_many(self, texts: list[str]) -> list[list[float]]:
        if not texts:
            return []
        return [self._normalize(vector.tolist() if hasattr(vector, "tolist") else list(vector)) for vector in self._model.embed(texts)]

    def embed(self, text: str) -> list[float]:
        return self.embed_many([text or " "])[0]

    def _normalize(self, vector: list[float]) -> list[float]:
        norm = math.sqrt(sum(value * value for value in vector))
        if norm == 0:
            return vector
        return [value / norm for value in vector]


class EmbeddingProviderFactory:
    @staticmethod
    def build() -> EmbeddingProvider:
        provider = settings.rag_embedding_provider.lower()
        if provider == "fastembed":
            try:
                return FastEmbedEmbeddingProvider(settings.rag_embedding_model)
            except Exception:
                if not settings.rag_allow_hashing_fallback:
                    raise
        return HashingEmbeddingProvider(dimension=settings.rag_embedding_dimension)


def _cosine_similarity(left: list[float], right: list[float]) -> float:
    if not left or not right or len(left) != len(right):
        return 0.0
    return sum(a * b for a, b in zip(left, right, strict=True))


class InMemoryVectorStore:
    backend_name = "memory"

    def __init__(self) -> None:
        self._items: dict[str, tuple[list[float], DocumentChunk]] = {}

    def upsert(self, chunks: list[DocumentChunk], embeddings: list[list[float]]) -> int:
        for chunk, embedding in zip(chunks, embeddings, strict=True):
            self._items[chunk.chunk_id] = (embedding, chunk)
        return len(chunks)

    def search(self, query_embedding: list[float], limit: int = 5) -> list[RetrievalResult]:
        scored = [
            RetrievalResult(
                chunk_id=chunk.chunk_id,
                document_id=chunk.document_id,
                text=chunk.text,
                score=_cosine_similarity(query_embedding, embedding),
                metadata=chunk.metadata,
            )
            for embedding, chunk in self._items.values()
        ]
        return sorted(scored, key=lambda item: item.score, reverse=True)[:limit]

    def get_chunks(self, document_ids: list[str], limit: int = 6) -> list[RetrievalResult]:
        results = [
            RetrievalResult(
                chunk_id=chunk.chunk_id,
                document_id=chunk.document_id,
                text=chunk.text,
                score=1.0,
                metadata=chunk.metadata,
            )
            for _, chunk in self._items.values()
            if chunk.document_id in document_ids
        ]
        return sorted(results, key=lambda item: item.metadata.get("chunk_index", 0))[:limit]


class QdrantVectorStore:
    backend_name = "qdrant"
    _client_cache: ClassVar[dict[str, Any]] = {}

    def __init__(self, collection_name: str, dimension: int, url: str | None = None, path: str | None = None) -> None:
        if QdrantClient is None or qdrant_models is None:
            raise RuntimeError("qdrant-client is not installed")
        self.collection_name = collection_name
        self.dimension = dimension
        if path:
            Path(path).mkdir(parents=True, exist_ok=True)
            cache_key = f"path:{Path(path).resolve()}"
            if cache_key not in self._client_cache:
                self._client_cache[cache_key] = QdrantClient(path=path)
            self.client = self._client_cache[cache_key]
        else:
            cache_key = f"url:{url or settings.qdrant_url}"
            if cache_key not in self._client_cache:
                self._client_cache[cache_key] = QdrantClient(url=url or settings.qdrant_url)
            self.client = self._client_cache[cache_key]

    def ensure_collection(self) -> None:
        assert qdrant_models is not None
        collections = self.client.get_collections().collections
        if any(collection.name == self.collection_name for collection in collections):
            return
        self.client.create_collection(
            collection_name=self.collection_name,
            vectors_config=qdrant_models.VectorParams(
                size=self.dimension,
                distance=qdrant_models.Distance.COSINE,
            ),
        )

    def upsert(self, chunks: list[DocumentChunk], embeddings: list[list[float]]) -> int:
        assert qdrant_models is not None
        self.ensure_collection()
        points = [
            qdrant_models.PointStruct(
                id=str(uuid.uuid5(uuid.NAMESPACE_URL, chunk.chunk_id)),
                vector=embedding,
                payload={
                    "chunk_id": chunk.chunk_id,
                    "document_id": chunk.document_id,
                    "text": chunk.text,
                    "metadata": chunk.metadata,
                },
            )
            for chunk, embedding in zip(chunks, embeddings, strict=True)
        ]
        self.client.upsert(collection_name=self.collection_name, points=points)
        return len(points)

    def search(self, query_embedding: list[float], limit: int = 5) -> list[RetrievalResult]:
        self.ensure_collection()
        points = self.client.search(
            collection_name=self.collection_name,
            query_vector=query_embedding,
            limit=limit,
            with_payload=True,
        )
        return [self._point_to_result(point, default_score=float(point.score)) for point in points]

    def get_chunks(self, document_ids: list[str], limit: int = 6) -> list[RetrievalResult]:
        assert qdrant_models is not None
        self.ensure_collection()
        points, _ = self.client.scroll(
            collection_name=self.collection_name,
            scroll_filter=qdrant_models.Filter(
                must=[
                    qdrant_models.FieldCondition(
                        key="document_id",
                        match=qdrant_models.MatchAny(any=document_ids),
                    )
                ]
            ),
            limit=limit,
            with_payload=True,
        )
        results = [self._point_to_result(point, default_score=1.0) for point in points]
        return sorted(results, key=lambda item: item.metadata.get("chunk_index", 0))

    def _point_to_result(self, point: Any, default_score: float) -> RetrievalResult:
        payload = point.payload or {}
        return RetrievalResult(
            chunk_id=str(payload.get("chunk_id", point.id)),
            document_id=str(payload.get("document_id", "")),
            text=str(payload.get("text", "")),
            score=default_score,
            metadata=dict(payload.get("metadata") or {}),
        )


class RagKnowledgeBaseService:
    def __init__(self) -> None:
        self.parser = DocumentParser()
        self.chunker = TextChunker(
            chunk_size=settings.rag_chunk_size,
            overlap=settings.rag_chunk_overlap,
        )
        self.embedding_provider = EmbeddingProviderFactory.build()
        self.vector_store = self._build_vector_store()

    def _build_vector_store(self) -> InMemoryVectorStore | QdrantVectorStore:
        if settings.rag_use_qdrant:
            try:
                return QdrantVectorStore(
                    collection_name=settings.rag_collection_name,
                    dimension=self.embedding_provider.dimension,
                    url=settings.qdrant_url,
                    path=settings.qdrant_path or None,
                )
            except Exception:
                if not settings.rag_allow_memory_fallback:
                    raise
        return InMemoryVectorStore()

    def ingest_document(
        self,
        content: bytes,
        filename: str,
        content_type: str | None = None,
        metadata: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        document = self.parser.parse(content, filename, content_type, metadata)
        chunks = self.chunker.split(document)
        embeddings = self.embedding_provider.embed_many([chunk.text for chunk in chunks])
        indexed_chunks = self.vector_store.upsert(chunks, embeddings) if chunks else 0
        return {
            "document_id": document.document_id,
            "filename": document.filename,
            "content_type": document.content_type,
            "size": document.metadata["byte_size"],
            "chunk_count": len(chunks),
            "indexed_chunks": indexed_chunks,
            "status": "indexed" if indexed_chunks else "empty",
            "metadata": {
                **document.metadata,
                "embedding_provider": self.embedding_provider.provider_name,
                "embedding_dimension": self.embedding_provider.dimension,
                "vector_store": self.vector_store.backend_name,
            },
        }

    def search(self, query: str, limit: int = 5, document_ids: list[str] | None = None) -> list[RetrievalResult]:
        query_embedding = self.embedding_provider.embed(query)
        results = self.vector_store.search(query_embedding, limit=limit)
        if document_ids:
            results = [result for result in results if result.document_id in document_ids]
        return results

    def get_document_context(self, document_ids: list[str], query: str = "", limit: int = 5) -> str:
        if not document_ids:
            return ""
        results = self.search(query, limit=limit, document_ids=document_ids) if query else []
        if not results and hasattr(self.vector_store, "get_chunks"):
            results = self.vector_store.get_chunks(document_ids, limit=limit)
        return "\n\n".join(
            f"资料：{result.metadata.get('filename', result.document_id)}\n{result.text}"
            for result in results
            if result.text.strip()
        )


rag_service = RagKnowledgeBaseService()
